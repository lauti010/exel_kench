import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from app import app


class ContactsTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        app.config.update(TESTING=True, EXCEL_PATH=Path(self.directory.name) / 'contacts.xlsx')
        self.client = app.test_client()

    def tearDown(self):
        self.directory.cleanup()

    def test_save_download_and_append(self):
        for name in ['Lautaro', '=1+1']:
            result = self.client.post('/api/contactos', json={
                'nombre': name, 'apellido': 'Pérez', 'gmail': 'Prueba@gmail.com'})
            self.assertEqual(result.status_code, 201)
        book = load_workbook(BytesIO(self.client.get('/descargar').data))
        self.assertEqual(list(book.active.values), [
            ('Nombre', 'Apellido', 'Gmail'),
            ('Lautaro', 'Pérez', 'prueba@gmail.com'),
            ('=1+1', 'Pérez', 'prueba@gmail.com')])
        self.assertEqual(book.active['A3'].data_type, 's')
        book.close()

    def test_reject_invalid_data_without_creating_file(self):
        for data in [None, [], {'nombre': '', 'apellido': 'Pérez', 'gmail': 'a@gmail.com'},
                     {'nombre': 'Ana', 'apellido': 'Pérez', 'gmail': 'a@example.com'}]:
            self.assertEqual(self.client.post('/api/contactos', json=data).status_code, 400)
        self.assertFalse(Path(app.config['EXCEL_PATH']).exists())

    def test_empty_download_has_headers(self):
        response = self.client.get('/descargar')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(list(book.active.values), [('Nombre', 'Apellido', 'Gmail')])
        book.close()


if __name__ == '__main__':
    unittest.main()
