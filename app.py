import os
import re
import tempfile
from io import BytesIO
from pathlib import Path
from threading import Lock

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024
app.config['EXCEL_PATH'] = Path(os.environ.get('EXCEL_PATH', 'data/contactos.xlsx'))
lock = Lock()


def workbook():
    path = Path(app.config['EXCEL_PATH'])
    if path.exists():
        return load_workbook(path)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Contactos'
    sheet.append(['Nombre', 'Apellido', 'Gmail'])
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='334155')
    for column, width in [('A', 28), ('B', 28), ('C', 42)]:
        sheet.column_dimensions[column].width = width
    return book


@app.get('/')
def index():
    return render_template('index.html')


@app.post('/api/contactos')
def save_contact():
    values = request.get_json(silent=True)
    if not isinstance(values, dict):
        return jsonify(error='Enviá los tres campos del formulario.'), 400
    fields = [values.get(key) for key in ('nombre', 'apellido', 'gmail')]
    if any(not isinstance(value, str) for value in fields):
        return jsonify(error='Completá nombre, apellido y Gmail.'), 400
    nombre, apellido, gmail = [value.strip() for value in fields]
    if any(not value or len(value) > 80 or any(ord(c) < 32 for c in value)
           for value in (nombre, apellido)):
        return jsonify(error='Nombre y apellido deben tener entre 1 y 80 caracteres.'), 400
    if len(gmail) > 254 or not re.fullmatch(r'[a-zA-Z0-9](?:[a-zA-Z0-9._+\-]*[a-zA-Z0-9])?@gmail\.com', gmail, re.I):
        return jsonify(error='Ingresá una dirección válida terminada en @gmail.com.'), 400
    with lock:
        book = workbook()
        sheet = book.active
        sheet.append([nombre, apellido, gmail.lower()])
        # Force text so submitted names cannot become Excel formulas.
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
        sheet.auto_filter.ref = sheet.dimensions
        path = Path(app.config['EXCEL_PATH'])
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as file:
                temporary = Path(file.name)
            book.save(temporary)
            os.replace(temporary, path)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
            book.close()
    return jsonify(message='¡Listo! Tus datos se guardaron en el Excel.'), 201


@app.get('/descargar')
def download():
    with lock:
        book = workbook()
        output = BytesIO()
        book.save(output)
        book.close()
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='contactos.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', '5000')))
